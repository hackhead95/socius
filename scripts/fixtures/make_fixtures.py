"""Generate the small SPSS fixtures used by tests/io (and the expected JSON pyreadstat reads from them).

Run from the repository root:
  /opt/oracle/bin/python scripts/fixtures/make_fixtures.py

Writes tests/io/fixtures/<name>.sav|.zsav and <name>.json (the pyreadstat reading of that file, see
dump_sav.py). A few fixtures are patched after writing to reproduce situations pyreadstat cannot
write directly:
  - weighted.sav:            header weight index set (pyreadstat has no WEIGHT BY support); the
                             expected weight variable name goes into the JSON as "weight_var".
  - ncases_unknown*.sav:     header case count set to -1 (the reader must count cases itself).
  - trailing_garbage.sav:    random bytes appended after the data (must be ignored).
"""
import collections
import collections.abc
import datetime as dt
import json
import os
import struct
import sys

import numpy as np
import pandas as pd
import pyreadstat

sys.path.insert(0, os.path.dirname(__file__))
from dump_sav import dump  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "tests", "io", "fixtures")

BN = "আমি বাংলায় কথা বলি"  # Bengali
HI = "नमस्ते दुनिया"  # Hindi
LATIN = "Crème brûlée, São Paulo, Łódź"
EMOJI = "Happy 😀 👍🏽 family 👨‍👩‍👧"


def exact_bytes(s, n, fill="x"):
    """Return s padded with `fill` to exactly n UTF-8 bytes (s itself must be shorter)."""
    b = s.encode("utf-8")
    assert len(b) <= n
    return s + fill * (n - len(b))


def main_frame():
    long600 = "A" * 255 + "B" * 255 + "C" * 90  # 600 bytes: checks the 255-byte segment layout
    # 598 bytes: 1 ASCII byte then 199 three-byte Bengali letters, so characters straddle both the
    # 255-byte very-long-string segment boundaries and the 8-byte chunk boundaries.
    split600 = "x" + "অ" * 199
    return pd.DataFrame({
        "id": [1.0, 2.0, 3.0, 4.0, 5.0, 6.0],
        "gender": [1.0, 2.0, 2.0, 9.0, 1.0, np.nan],
        "likert": [1.0, 5.0, 3.0, 8.0, -1.0, 4.0],
        "income": [52000.5, -3.0, 0.0, 1234567.25, np.nan, 250.75],
        "score": [0.125, 99.0, -99.0, 150.0, 151.0, -100.0],
        "a_very_long_variable_name_for_testing": [1.0, 2.0, 3.0, 4.0, 5.0, 6.0],
        "বয়স": [23.0, 35.0, 41.0, 19.0, 60.0, 33.0],
        "s1": ["a", "b", "", "z", "a", " "],
        "s8": ["abcdefgh", "12345678", "x", "", "short", "é"],
        "s9": ["abcdefghi", "123456789", "y", "", "nine", "ü"],
        "s255": [exact_bytes("start", 255), "mid", "", "end", "z", "q"],
        "s256": [exact_bytes("start", 256, "y"), "m", "", "e", "z", "q"],
        "s600": [long600, "short", "", split600, "z", "q"],
        "sex_str": ["m", "f", "f", "x", "m", "f"],
        "region": ["North region", "South region", "East", "West", "North region", "unknown value"],
        "text_bn": [BN, HI, LATIN, EMOJI, "", "plain"],
        "d": [dt.date(2020, 1, 2), dt.date(1999, 12, 31), None, dt.date(1582, 10, 15), dt.date(2024, 2, 29), dt.date(1970, 1, 1)],
        "ts": [dt.datetime(2020, 1, 2, 3, 4, 5), None, dt.datetime(1970, 1, 1), dt.datetime(2001, 9, 9, 1, 46, 40), dt.datetime(1900, 1, 1), dt.datetime(2030, 12, 31, 23, 59, 59)],
        "t": [dt.time(1, 2, 3), None, dt.time(23, 59, 59), dt.time(0, 0, 0), dt.time(12, 0, 0), dt.time(6, 30, 15)],
    })


def main_kwargs():
    return dict(
        file_label="Socius test fixture: বাংলা",
        column_labels={
            "id": "Respondent ID",
            "gender": "Respondent gender",
            "likert": "Agree with statement (1-5)",
            "income": "Household income in rupees",
            "score": "Test score",
            "a_very_long_variable_name_for_testing": "A variable with a long name",
            "বয়স": "বয়স (বছর)",
            "text_bn": "Open-ended answer: " + HI,
            "d": "Interview date",
            "ts": "Interview start",
            "t": "Interview duration",
            "region": "Region of residence",
        },
        note=["First document line.", "Second line with Unicode: " + BN],
        variable_value_labels={
            "gender": {1: "Male", 2: "Female", 9: "No answer"},
            "likert": {1: "Strongly disagree", 2: "Disagree", 3: "Neutral", 4: "Agree", 5: "Strongly agree", 8: "Don't know", -1: "Refused"},
            "score": {-99: "Not tested", 0.125: "One eighth"},
            "sex_str": {"m": "Male", "f": "Female", "x": "Other/unknown"},
            "region": {"North region": "উত্তর", "South region": "दक्षिण"},
            "বয়স": {23: "তেইশ"},
        },
        missing_ranges={
            "gender": [9],
            "likert": [{"lo": 8, "hi": 9}, -1],
            "income": [{"lo": float("-inf"), "hi": 0}],
            "score": [-99, 150, 151],
            "sex_str": ["x"],
            "region": ["unknown"],
        },
        variable_measure={
            "id": "scale",
            "gender": "nominal",
            "likert": "ordinal",
            "income": "scale",
            "sex_str": "nominal",
            "region": "nominal",
        },
        variable_display_width={"id": 6, "income": 12, "text_bn": 30},
        variable_format={"id": "F6.0", "income": "DOLLAR12.2", "score": "F8.3", "gender": "F1.0"},
    )


def write_with_json(df, name, **kwargs):
    path = os.path.join(OUT, name)
    pyreadstat.write_sav(df, path, **kwargs)
    return path


def save_json(path, extra=None):
    data = dump(path)
    if extra:
        data.update(extra)
    with open(os.path.splitext(path)[0] + ".json", "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1, allow_nan=False)


def patch_int32(path, offset, value):
    with open(path, "r+b") as f:
        f.seek(64)
        layout = struct.unpack("<i", f.read(4))[0]
        endian = "<" if layout in (2, 3) else ">"
        f.seek(offset)
        f.write(struct.pack(endian + "i", value))


def make_xlsx(path):
    """A two-sheet workbook with the cell types Excel users produce (read by tests/io/xlsx.test.ts)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Survey"
    ws.append(["id", "name", "score", "joined", "started", "duration", "agree", "notes", "Age (years)", "বয়স", "empty"])
    rows = [
        [1, "Asha", 3.5, dt.date(2021, 3, 4), dt.datetime(2021, 3, 4, 9, 30, 0), dt.time(0, 45, 0), True, 12, 34, 20, None],
        [2, "Bikram", None, dt.date(1999, 12, 31), dt.datetime(2000, 1, 1, 0, 0, 1), dt.time(1, 2, 3), False, "see note", 41, 21, None],
        [3, "Chandra", 10, None, None, None, True, None, None, 22, None],
        [4, "দীপা", -2.25, dt.date(2024, 2, 29), dt.datetime(2024, 2, 29, 23, 59, 59), dt.time(12, 0, 0), None, 7.5, 29, 23, None],
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in row:
            c.number_format = "yyyy-mm-dd hh:mm:ss"
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row:
            c.number_format = "hh:mm:ss"
    other = wb.create_sheet("Codes")
    other.append(["code", "meaning"])
    other.append([1, "Yes"])
    other.append([2, "No"])
    wb.save(path)


def make_spss_written():
    """Files written by IBM's SPSS I/O library (libspssdio, bundled with savReaderWriter): what SPSS
    itself produces. Expected readings come from the same library (spssio_dump.py). Skipped with a
    note when savReaderWriter is not installed; the committed files stay valid."""
    import subprocess

    for n in ("Iterable", "Mapping", "MutableMapping", "Sequence", "Callable"):
        if not hasattr(collections, n):
            setattr(collections, n, getattr(collections.abc, n))
    try:
        import savReaderWriter as srw
    except ImportError:
        print("savReaderWriter not installed: SPSS-written fixtures not regenerated")
        return
    var_names = [b"id", b"gender", b"income", b"region", b"essay", b"q1_a_long_variable_name", b"when", b"wt"]
    var_types = {b"id": 0, b"gender": 0, b"income": 0, b"region": 12, b"essay": 700, b"q1_a_long_variable_name": 0, b"when": 0, b"wt": 0}
    essay = ("আমরা " * 50 + "end").encode("utf-8")  # 653 bytes, spans 3 very-long-string segments
    records = [
        [1, 1, 52000.5, "North".encode(), essay, 3, 13797302400.0, 1.0],
        [2, 2, None, "South east".encode(), b"short", 5, None, 2.5],
        [3, 9, -1, "".encode(), "Crème brûlée".encode(), 1, 12219379200.0, 0.5],
    ]
    kwargs = dict(
        valueLabels={b"gender": {1: b"Male", 2: b"Female", 9: b"No answer"}, b"region": {b"North": "উত্তর".encode()}},
        varLabels={b"id": b"Respondent", b"gender": b"Gender", b"essay": "Open answer (বাংলা)".encode(), b"q1_a_long_variable_name": b"Question 1"},
        formats={b"id": b"F4.0", b"income": b"DOLLAR12.2", b"when": b"ADATE10", b"region": b"A12", b"essay": b"A700"},
        missingValues={b"gender": {"values": [9]}, b"income": {"lower": -1.7976931348623155e308, "upper": 0}, b"region": {"values": [b"none"]}},
        measureLevels={b"gender": b"nominal", b"q1_a_long_variable_name": b"ordinal", b"income": b"ratio", b"region": b"nominal", b"essay": b"nominal", b"id": b"ratio", b"when": b"ratio", b"wt": b"ratio"},
        columnWidths={b"essay": 40, b"region": 12},
        alignments={b"id": b"left"},
        varRoles={b"q1_a_long_variable_name": b"target", b"wt": b"none"},
        varAttributes={b"income": {b"Source": b"Q12"}},
        fileLabel="SPSS I/O fixture".encode(),
        multRespDefs={b"$mr": {b"setType": b"D", b"label": b"Multi", b"varNames": [b"gender", b"q1_a_long_variable_name"], b"countedValue": b"1"}},
        caseWeightVar=b"wt",
        ioUtf8=True,
    )
    script = os.path.join(os.path.dirname(__file__), "spssio_dump.py")
    for name in ("spss_io_bytecode.sav", "spss_io_zlib.zsav"):
        path = os.path.join(OUT, name)
        with srw.SavWriter(path, var_names, var_types, **kwargs) as w:
            for r in records:
                w.writerow(r)
        out = subprocess.run([sys.executable, script, path], check=True, capture_output=True).stdout
        with open(os.path.splitext(path)[0] + ".spssio.json", "wb") as f:
            f.write(out)
        save_json(path, {"weight_var": "wt"})


def main():
    os.makedirs(OUT, exist_ok=True)
    for fn in os.listdir(OUT):
        if fn.endswith((".sav", ".zsav", ".json", ".xlsx")):  # includes *.spssio.json
            os.remove(os.path.join(OUT, fn))

    df = main_frame()
    kw = main_kwargs()
    save_json(write_with_json(df, "main_uncompressed.sav", **kw))
    save_json(write_with_json(df, "main_bytecode.sav", row_compress=True, **kw))
    save_json(write_with_json(df, "main_zlib.zsav", compress=True, **kw))

    # Zero cases, still with a full dictionary.
    empty = df.iloc[0:0].copy()
    save_json(write_with_json(empty, "empty_bytecode.sav", row_compress=True, **kw))
    save_json(write_with_json(empty, "empty_zlib.zsav", compress=True, **kw))

    # Weight variable (pyreadstat cannot set it, so patch the header's weight index; dictionary
    # index 1 is the first variable record).
    wdf = pd.DataFrame({"wt": [1.0, 2.5, 0.5, 1.0], "y": [1.0, 2.0, 3.0, 4.0]})
    p = write_with_json(wdf, "weighted.sav", row_compress=True)
    patch_int32(p, 76, 1)
    save_json(p, {"weight_var": "wt"})

    # Case count unknown (-1): the reader must count cases by reading the data.
    ndf = pd.DataFrame({"n": [1.0, 2.0, np.nan, 1e10, -0.5], "s": ["a", "bb", "", "dddddddddd", "e"]})
    for name, kwargs in (("ncases_unknown.sav", {}), ("ncases_unknown_bytecode.sav", {"row_compress": True})):
        p = write_with_json(ndf, name, **kwargs)
        # Dump before patching: readstat itself returns 0 rows for an uncompressed file with an
        # unknown case count, so the unpatched reading is the ground truth.
        save_json(p, {"header_ncases": -1})
        patch_int32(p, 80, -1)

    # Trailing garbage after the data (the case count is known, so the reader stops in time).
    p = write_with_json(ndf, "trailing_garbage.sav", row_compress=True)
    save_json(p)
    with open(p, "ab") as f:
        f.write(bytes(range(256)) * 3)

    make_xlsx(os.path.join(OUT, "workbook.xlsx"))
    make_spss_written()

    total = sum(os.path.getsize(os.path.join(OUT, f)) for f in os.listdir(OUT))
    print(f"fixtures written to {OUT}: {len(os.listdir(OUT))} files, {total} bytes")


if __name__ == "__main__":
    main()
